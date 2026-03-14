import subprocess
from datetime import datetime, date, timedelta
from typing import Optional
import sys

# ---------------------------------------------------------------------------
# Auto-instalación de dependencias
# ---------------------------------------------------------------------------
def instalar_si_falta(paquete: str, extra_cmd: list = None) -> None:
    try:
        __import__(paquete)
    except ImportError:
        print(f"'{paquete}' no encontrado. Instalando...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", paquete])
        if extra_cmd:
            subprocess.check_call(extra_cmd)
        print(f"'{paquete}' instalado correctamente.")

instalar_si_falta("openpyxl")
instalar_si_falta(
    "playwright",
    extra_cmd=[sys.executable, "-m", "playwright", "install", "chromium"],
)

import openpyxl
from playwright.sync_api import sync_playwright

# ---------------------------------------------------------------------------
# CONFIGURACIÓN — ajusta estos valores
# NOTA: Si tu certificado está en un token hardware (tarjeta inteligente, USB),
#       Playwright no puede leerlo directamente. En ese caso necesitarías
#       exportar el certificado a un archivo .pfx primero.
# ---------------------------------------------------------------------------
EXCEL_PATH = r"C:\Users\Usuario\OneDrive\PROYECTOS DIGITALES\TRABAJOS PERSONALIZADOS\GESTION.ES\AUTOMATIZACIONES SELENIUM\ACCESO_DEHU_ULTIMO.xlsm"
CERT_PATH  = r"C:\ruta\al\certificado.pfx"  # .pfx o .p12
CERT_PASS  = "contraseña_certificado"
NIF        = "60560345B"

# ---------------------------------------------------------------------------
# UTILIDADES
# ---------------------------------------------------------------------------
def get_last_row(ws: openpyxl.worksheet.worksheet.Worksheet) -> int:
    """Devuelve el número de la última fila con datos en la columna A."""
    last = 1
    for row in ws.iter_rows(min_col=1, max_col=1):
        for cell in row:
            if cell.value is not None:
                last = cell.row
    return last


def parse_date(value: str) -> Optional[date]:
    """Convierte string de fecha española (dd/mm/yyyy) a objeto date."""
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(value.strip(), fmt).date()
        except (ValueError, AttributeError):
            continue
    return None

# ---------------------------------------------------------------------------
# SCRIPTS JAVASCRIPT
# En Playwright, page.evaluate() espera una función JS: "() => { ... }"
# ---------------------------------------------------------------------------
def _js_rows_count() -> str:
    return """() => {
        const t = document.querySelector('#tablaNotificacionesPendientes');
        if (!t || !t.shadowRoot) return 0;
        const tb = t.shadowRoot.querySelector('table[data-name="DntTable"]');
        if (!tb) return 0;
        return tb.querySelectorAll('tr.dnt-table__row.dnt-table__row--grouped').length;
    }"""


def _js_cell(row_idx: int, col_idx: int) -> str:
    return f"""() => {{
        const t  = document.querySelector('#tablaNotificacionesPendientes');
        const rows = t.shadowRoot
            .querySelector('table[data-name="DntTable"]')
            .querySelectorAll('tr.dnt-table__row.dnt-table__row--grouped');
        return rows[{row_idx}].querySelectorAll('td')[{col_idx}].innerText.trim();
    }}"""


def _js_objeto(row_idx: int) -> str:
    return f"""() => {{
        const t  = document.querySelector('#tablaNotificacionesPendientes');
        const rows = t.shadowRoot
            .querySelector('table[data-name="DntTable"]')
            .querySelectorAll('tr.dnt-table__row.dnt-table__row--grouped');
        const td = rows[{row_idx}].querySelectorAll('td')[3].querySelector('div');
        return td?.getAttribute('data-original-text-concept') ?? td?.innerText.trim() ?? '';
    }}"""


def _js_fecha_caducidad(t_idx: int) -> str:
    """
    Accede directamente a las filas de tbody para leer la fecha de caducidad.
    t_idx = 2*i - 1  (fórmula corregida; el código anterior era incorrecto para i>=4)
    """
    return f"""() => {{
        const t = document.querySelector('#tablaNotificacionesPendientes');
        const rows = t.shadowRoot.querySelectorAll('tbody tr');
        if (!rows[{t_idx}]) return '';
        return rows[{t_idx}].querySelectorAll('td')[1].innerText.trim();
    }}"""

# ---------------------------------------------------------------------------
# FUNCIÓN PRINCIPAL: consulta()
# ---------------------------------------------------------------------------
def consulta(nif: str = NIF, conteo: int = 1, hoja2_fila: int = 2) -> str:
    """
    Accede a la DEHú con Playwright (headless, sin AutoIT) y extrae
    notificaciones pendientes al Excel.

    Parámetros
    ----------
    nif        : NIF a consultar
    conteo     : contador de ciclo
    hoja2_fila : fila de Hoja2 donde escribir el estado del proceso
    """
    wb  = openpyxl.load_workbook(EXCEL_PATH, keep_vba=True)
    ws1 = wb.worksheets[0]   # Hoja1 — datos principales
    ws2 = wb.worksheets[1]   # Hoja2 — estado del proceso
    estado = "ERROR"

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        context = browser.new_context(
            client_certificates=[{
                "origin": "https://dehu.redsara.es",
                "pfxPath": CERT_PATH,
                "passphrase": CERT_PASS,
            }]
        )
        page = context.new_page()

        try:
            # ---- Abrir DEHú y pulsar botón de acceso ----
            page.goto("https://dehu.redsara.es", wait_until="domcontentloaded")
            # Esperar a que el web component cargue su shadow DOM antes de clicar
            page.wait_for_selector("dnt-button.access-btn", timeout=15000)
            page.evaluate(
                "() => {"
                "  const btn = document.querySelector('dnt-button.access-btn');"
                "  const inner = btn.shadowRoot ? btn.shadowRoot.querySelector('button') : btn;"
                "  inner.click();"
                "}"
            )
            page.wait_for_load_state("networkidle")

            # ---- Seleccionar certificado digital ----
            page.wait_for_selector(
                "xpath=//*[@id='ID_main']/div[2]/div/div/div/article[2]/div[4]/button",
                timeout=15000,
            )
            page.click(
                "xpath=//*[@id='ID_main']/div[2]/div/div/div/article[2]/div[4]/button/span[1]"
            )
            page.wait_for_load_state("networkidle")

            # ---- Ir a notificaciones pendientes ----
            page.goto("https://dehu.redsara.es/es/notifications", wait_until="networkidle")
            current_url = page.url
            if "security" in current_url.upper():
                page.goto("https://dehu.redsara.es/es/notifications", wait_until="networkidle")
                current_url = page.url

            # ---- Verificar acceso correcto ----
            last_row = get_last_row(ws1) + 1
            if current_url != "https://dehu.redsara.es/es/notifications":
                ws1.cell(row=last_row, column=1).value = datetime.now()
                ws1.cell(row=last_row, column=2).value = nif
                ws1.cell(row=last_row, column=4).value = "ERROR REVISAR CERTIFICADO"
                ws2.cell(row=hoja2_fila, column=3).value = "Error al acceder con el certificado"
                wb.save(EXCEL_PATH)
                return "ERROR"

            # ---- Extraer NIF y Razón Social del encabezado ----
            texto = page.locator("xpath=//*[@id='pane-0']/p[1]").inner_text().upper()
            prefijo = "LAS NOTIFICACIONES PENDIENTES QUE SE MUESTRAN SE HAN EMITIDO A "
            texto_limpio = texto.replace(prefijo, "")
            idx_nif = texto_limpio.find("NIF")
            q_nif = texto_limpio[idx_nif + 4: idx_nif + 13].strip()

            if nif.upper().strip() != q_nif:
                ws1.cell(row=last_row, column=1).value = datetime.now()
                ws1.cell(row=last_row, column=2).value = nif
                ws1.cell(row=last_row, column=4).value = "ERROR REVISAR CERTIFICADO"
                ws2.cell(row=hoja2_fila, column=3).value = "Error al acceder con el certificado"
                wb.save(EXCEL_PATH)
                return "ERROR"

            idx_con = texto_limpio.find("CON EL NIF")
            q_rz = texto_limpio[:idx_con].strip()

            # ---- Contar notificaciones ----
            no_notif_elem = page.evaluate(_js_rows_count())

            # ---- CASO: Sin notificaciones ----
            if not no_notif_elem or no_notif_elem <= 0:
                ws1.cell(row=last_row, column=1).value = datetime.now()
                ws1.cell(row=last_row, column=2).value = q_nif
                ws1.cell(row=last_row, column=3).value = q_rz
                txt_sin = page.evaluate(
                    "() => document.querySelector('app-notifications-list p.ng-star-inserted')"
                    "?.innerText.trim() ?? ''"
                )
                if "no dispones de notificaciones" in txt_sin.lower():
                    ws1.cell(row=last_row, column=4).value = (
                        "No se encontraron resultados para los filtros seleccionados"
                    )
                estado = "Terminado"

            # ---- CASO: Existen notificaciones ----
            else:
                txt_info = page.evaluate(
                    "() => {"
                    "  const p = document.querySelector('#pagination-PendingNotifications');"
                    "  if (!p || !p.shadowRoot) return '';"
                    "  const el = p.shadowRoot.querySelector('.dnt-pagination-total');"
                    "  return el ? el.textContent.trim() : '';"
                    "}"
                )
                pag = 1
                if txt_info:
                    try:
                        total = int(txt_info.split(" de ")[1].split(" ")[0].strip())
                        pag = (total + 9) // 10
                    except (IndexError, ValueError):
                        pag = 1

                for q in range(1, pag + 1):
                    if q > 1:
                        page.evaluate(
                            "() => {"
                            "  const pag = document.querySelector('#pagination-PendingNotifications');"
                            "  const realBtn = pag?.shadowRoot"
                            "    ?.querySelector('dnt-pagination-next')"
                            "    ?.querySelector('dnt-button')"
                            "    ?.shadowRoot?.querySelector('button');"
                            "  realBtn?.click();"
                            "}"
                        )
                        page.wait_for_load_state("networkidle")
                        page.evaluate("() => window.scrollBy(0, 200)")
                        no_notif_elem = page.evaluate(_js_rows_count())

                    num_notif = no_notif_elem
                    page.evaluate("() => window.scrollBy(0, 200)")

                    for i in range(1, num_notif + 1):
                        if not (i == 1 and conteo > 1):
                            last_row = get_last_row(ws1) + 1
                        idx = i - 1   # índice 0-based para JS

                        ws1.cell(row=last_row, column=1).value  = datetime.now()
                        ws1.cell(row=last_row, column=2).value  = q_nif
                        ws1.cell(row=last_row, column=10).value = q_nif
                        ws1.cell(row=last_row, column=3).value  = q_rz

                        # Organismo (col 4, índice JS 4)
                        ws1.cell(row=last_row, column=4).value = page.evaluate(_js_cell(idx, 4))

                        # Objeto (col 5, índice JS 3)
                        ws1.cell(row=last_row, column=5).value = page.evaluate(_js_objeto(idx))

                        # Fecha Disposición (col 6, índice JS 6)
                        f1_str = page.evaluate(_js_cell(idx, 6))
                        f1 = parse_date(f1_str) if f1_str else None
                        ws1.cell(row=last_row, column=6).value = f1 if f1 else f1_str

                        # Identificador (col 7, índice JS 2) — guardado como texto con '
                        identificador = page.evaluate(_js_cell(idx, 2))
                        ws1.cell(row=last_row, column=7).value = "'" + str(identificador or "")

                        # Tipo (col 8, índice JS 7)
                        ws1.cell(row=last_row, column=8).value = page.evaluate(_js_cell(idx, 7))

                        # Fecha Caducidad (col 9)
                        # CORRECCIÓN: la fórmula original fallaba para i pares >= 4.
                        # La fórmula correcta es siempre t_idx = 2*i - 1:
                        #   i=1 → 1, i=2 → 3, i=3 → 5, i=4 → 7, ...
                        t_idx = 2 * i - 1
                        f2_str = page.evaluate(_js_fecha_caducidad(t_idx))
                        if f2_str:
                            f2_str = f2_str[-10:]   # últimos 10 caracteres (Right 10)
                        f2 = parse_date(f2_str) if f2_str else None
                        ws1.cell(row=last_row, column=9).value = f2 if f2 else f2_str

                estado = "Terminado"

            wb.save(EXCEL_PATH)

        except Exception as exc:
            print(f"[ERROR] {exc}")
            estado = "ERROR"
            try:
                wb.save(EXCEL_PATH)
            except Exception:
                pass
        finally:
            context.close()
            browser.close()

    return estado

# ---------------------------------------------------------------------------
# FUNCIÓN SECUNDARIA: historico()
# ---------------------------------------------------------------------------
def historico() -> None:
    """
    Mueve a Hoja3 las filas de Hoja1 con fecha de hace 31 días
    y las elimina de Hoja1.
    """
    print(f"Inicia: {datetime.now()}")
    wb  = openpyxl.load_workbook(EXCEL_PATH, keep_vba=True)
    ws1 = wb.worksheets[0]   # Hoja1
    ws3 = wb.worksheets[2]   # Hoja3

    # Detectar el número real de columnas para no perder datos
    max_col = ws1.max_column

    date_ref = date.today() - timedelta(days=31)
    filas_a_mover = []

    for row in ws1.iter_rows(min_row=2):
        cell_fecha = row[0].value
        if cell_fecha is None:
            continue
        if isinstance(cell_fecha, datetime):
            fecha = cell_fecha.date()
        elif isinstance(cell_fecha, date):
            fecha = cell_fecha
        else:
            fecha = parse_date(str(cell_fecha))
        if fecha == date_ref:
            filas_a_mover.append(row[0].row)

    for fila_num in filas_a_mover:
        ultima_h3 = get_last_row(ws3) + 1
        for col in range(1, max_col + 1):
            ws3.cell(row=ultima_h3, column=col).value = ws1.cell(row=fila_num, column=col).value

    for fila_num in sorted(filas_a_mover, reverse=True):
        ws1.delete_rows(fila_num)

    wb.save(EXCEL_PATH)
    print(f"Termina: {datetime.now()}")

# ---------------------------------------------------------------------------
# PUNTO DE ENTRADA
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    resultado = consulta(nif=NIF)
    print(f"Estado consulta: {resultado}")
    historico()
